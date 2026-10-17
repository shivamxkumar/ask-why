import io

import pytest

from app.parsing import ParseError, SUPPORTED_EXTENSIONS, UnsupportedFileError, parse_file
from conftest import FILES


def _parse(name, sid="S1"):
    return parse_file((FILES / name).read_bytes(), name, sid)


def test_markdown_units():
    src, units = _parse("meridian_project_brief.md")
    assert src.kind == "md" and src.unit_count == len(units)
    u = {x.id: x for x in units}
    assert u["S1.001"].kind == "heading" and u["S1.001"].label == "line 1"
    assert u["S1.002"].text.startswith("SYNTHETIC SAMPLE")  # '> ' stripped
    goal = u["S1.007"]
    assert goal.kind == "list_item" and goal.text == "Public launch on 15 March 2027, ahead of the spring maintenance season."
    assert goal.location.line == 13 and goal.location.paragraph == 7 and goal.label == "line 13"
    assert goal.section == "2. Goals"
    assert [x.id for x in units] == [f"S1.{i:03d}" for i in range(1, len(units) + 1)]


def test_text_steps_are_units():
    _, units = _parse("onboarding_workflow.txt", "S2")
    steps = [x for x in units if x.text.startswith("Step ")]
    assert len(steps) == 5 and all(x.kind == "list_item" for x in steps)
    assert steps[0].label == "line 6" and steps[0].id == "S2.003"


def test_pdf_units():
    src, units = _parse("engineering_delivery_plan.pdf", "S3")
    assert src.page_count == 3 and src.summary == f"3 pages · {len(units)} passages"
    u = {x.id: x for x in units}
    assert u["S3.p1.07"].text == "General availability is targeted for 2 May 2027."
    assert u["S3.p1.07"].label == "p.1 ¶7" and u["S3.p1.07"].location.page == 1
    assert u["S3.p3.04"].text == "No fallback routing provider has been evaluated yet."
    assert u["S3.p2.01"].kind == "heading" and u["S3.p2.02"].section == "Requirements"
    assert not any("· page" in x.text for x in units)  # running footer removed
    assert "All details are fictional." in u["S3.p1.02"].text  # line join


def test_csv_units_and_profile():
    src, units = _parse("pilot_accounts.csv", "S4")
    rows = [x for x in units if x.kind == "table_row"]
    assert src.row_count == 124 and len(rows) == 124 and src.summary == "124 rows × 8 columns"
    r9 = next(x for x in rows if x.id == "S4.r0009")
    assert r9.label == "row 9" and r9.location.row == 9 and "| technicians:  |" in r9.text
    prof = src.profiles[0]
    assert prof.rows == 124 and prof.columns == 8 and len(prof.duplicate_rows) == 4
    cols = {c.name: c for c in prof.column_profiles}
    assert cols["technicians"].missing == 7
    assert any("NE" in g and "Northeast" in g and "North East" in g for g in cols["region"].category_variants)
    stat = {x.id: x for x in units if x.kind == "dataset_stat"}
    assert "4 rows exactly duplicate an earlier row" in stat["S4.stat.dataset"].text
    assert stat["S4.stat.dataset"].label == "dataset profile"
    assert "9999" in stat["S4.stat.col.jobs-per-week"].text and "sentinel" in stat["S4.stat.col.jobs-per-week"].text
    assert "rows 19 (11)" in stat["S4.stat.col.nps"].text
    assert stat["S4.stat.col.region"].label == "column region"
    assert "Counts: 2026-03 (22)" in stat["S4.stat.col.signup-month"].text
    assert "plan_status = paid but monthly_fee_usd = 0" in stat["S4.stat.dataset"].text


def test_csv_profile_text_is_deterministic():
    _, a = _parse("pilot_accounts.csv", "S4")
    _, b = _parse("pilot_accounts.csv", "S4")
    assert [x.text for x in a] == [x.text for x in b]


def test_xlsx_units():
    src, units = _parse("finance_budget.xlsx", "S5")
    assert src.sheet_names == ["Budget", "Assumptions"]
    u = {x.id: x for x in units}
    total = u["S5.budget.r0008"]
    assert total.label == "Budget · row 8" and total.location.sheet == "Budget" and total.location.row == 8
    assert total.text == "Category: Total | Amount (USD): 1397000 | Notes: Finance working total, SYNTHETIC SAMPLE"
    assert u["S5.assumptions.r0001"].text == "Assumption: Launch month | Value: April 2027 | Owner: Finance"
    assert u["S5.budget.stat.dataset"].label == "Budget · profile"
    assert "S5.assumptions.stat.col.owner" in u


def test_docx_units():
    import docx

    d = docx.Document()
    d.add_heading("Plan", level=1)
    d.add_paragraph("We assume 40% adoption by Q3 2027.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Cost"
    t.cell(1, 0).text, t.cell(1, 1).text = "Hosting", "$5,000"
    buf = io.BytesIO()
    d.save(buf)
    src, units = parse_file(buf.getvalue(), "plan.docx", "S9")
    assert [x.id for x in units] == ["S9.001", "S9.002", "S9.003", "S9.004"]
    assert units[0].kind == "heading" and units[1].label == "¶2" and units[1].section == "Plan"
    assert units[3].label == "table 1 · row 2" and units[3].text == "Hosting | $5,000"


@pytest.mark.parametrize("name", ["virus.exe", "old.xls", "old.doc", "noext"])
def test_unsupported_rejected(name):
    with pytest.raises(UnsupportedFileError):
        parse_file(b"hello", name, "S1")


def test_empty_and_binary_rejected():
    with pytest.raises(ParseError):
        parse_file(b"", "a.txt", "S1")
    with pytest.raises(ParseError):
        parse_file(b"\x00\x01\x02binary", "a.txt", "S1")
    with pytest.raises(ParseError):
        parse_file(b"x" * (15 * 1024 * 1024 + 1), "big.txt", "S1")


def test_latin1_fallback_warns():
    src, units = parse_file("Café plan costs €5".encode("latin-1", "replace"), "a.txt", "S1")
    assert units and src.warnings


def test_scanned_pdf_warning():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    src, units = parse_file(doc.tobytes(), "scan.pdf", "S1")
    assert units == [] and any("OCR" in w for w in src.warnings)


def test_supported_list():
    assert set(SUPPORTED_EXTENSIONS) == {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".xlsx"}
